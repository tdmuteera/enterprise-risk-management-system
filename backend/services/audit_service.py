"""
Immutable audit trail service for forensic logging.

Features:
  - Insert-only RiskHistory records (never update/delete)
  - JSON diffs capturing only changed fields
  - IP address and user-agent from FastAPI requests
  - Audit trail retrieval and filtering
  - Excel export for forensic reports
"""

from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional
from uuid import UUID

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, desc
from starlette.requests import Request

from models.risk import RiskHistory
from utils.json_utils import compute_json_diff


async def log_action(
    db: AsyncSession,
    risk_id: UUID,
    user_id: UUID,
    action: str,
    previous_values: Optional[Dict[str, Any]] = None,
    new_values: Optional[Dict[str, Any]] = None,
    request: Optional[Request] = None,
) -> RiskHistory:
    """
    Create immutable audit trail entry for a risk action.

    This function ONLY INSERTS - never updates or deletes RiskHistory.
    Each action is permanently recorded with IP address, user-agent, and JSON diffs.

    Args:
        db: SQLAlchemy async session
        risk_id: UUID of risk being modified
        user_id: UUID of user performing action
        action: Action type (created, updated, status_changed, score_changed, deleted)
        previous_values: Dict of old values (optional, only changed fields)
        new_values: Dict of new values (optional, only changed fields)
        request: FastAPI Request object (optional, for IP and user-agent)

    Returns:
        RiskHistory record (inserted)

    Example:
        entry = await log_action(
            db=db,
            risk_id=risk_id,
            user_id=user_id,
            action="status_changed",
            previous_values={"status": "Open"},
            new_values={"status": "Mitigated"},
            request=request
        )
    """
    # Extract IP address and user-agent (proxy-aware)
    ip_address = None
    user_agent = None

    if request:
        # Prefer X-Forwarded-For when behind a reverse proxy
        xff = request.headers.get("x-forwarded-for")
        if xff:
            ip_address = xff.split(",")[0].strip() or None
        else:
            ip_address = request.client.host if request.client else None

        user_agent = request.headers.get("user-agent")

    previous_values = previous_values or {}
    new_values = new_values or {}

    # Store only changed fields (field-level old/new), keeping creates/deletes readable.
    stored_previous: Dict[str, Any]
    stored_new: Dict[str, Any]

    if action == "created":
        stored_previous = {}
        stored_new = new_values
    elif action == "deleted":
        stored_previous = previous_values
        stored_new = {}
    else:
        diff = compute_json_diff(previous_values, new_values)
        stored_previous = {}
        stored_new = {}

        for key, change in diff.get("modified", {}).items():
            stored_previous[key] = change.get("from")
            stored_new[key] = change.get("to")

        for key, value in diff.get("added", {}).items():
            stored_previous[key] = None
            stored_new[key] = value

        for key, value in diff.get("removed", {}).items():
            stored_previous[key] = value
            stored_new[key] = None

    # Create immutable history entry
    history_entry = RiskHistory(
        risk_id=risk_id,
        action=action,
        changed_by_id=user_id,
        changed_at=datetime.utcnow(),
        previous_values=stored_previous,
        new_values=stored_new,
        ip_address=ip_address,
        user_agent=user_agent,
    )

    db.add(history_entry)
    await db.flush()  # Get ID without committing

    return history_entry


async def get_audit_trail(
    db: AsyncSession,
    risk_id: UUID,
    limit: int = 100,
    offset: int = 0,
) -> List[RiskHistory]:
    """
    Retrieve audit trail for a specific risk.

    Returns all RiskHistory entries for a risk, ordered by newest first.

    Args:
        db: SQLAlchemy async session
        risk_id: UUID of risk
        limit: Max records to return (default 100)
        offset: Skip N records for pagination

    Returns:
        List of RiskHistory entries (newest first)
    """
    stmt = (
        select(RiskHistory)
        .where(RiskHistory.risk_id == risk_id)
        .order_by(desc(RiskHistory.changed_at))
        .limit(limit)
        .offset(offset)
    )

    result = await db.execute(stmt)
    return result.scalars().all()


async def get_user_audit_log(
    db: AsyncSession,
    user_id: UUID,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
    limit: int = 500,
) -> List[RiskHistory]:
    """
    Retrieve all actions performed by a specific user within a date range.

    Useful for user activity monitoring and forensic investigation.

    Args:
        db: SQLAlchemy async session
        user_id: UUID of user
        start_date: Filter from this date (default: 30 days ago)
        end_date: Filter to this date (default: now)
        limit: Max records to return

    Returns:
        List of RiskHistory entries for the user
    """
    if not start_date:
        start_date = datetime.utcnow() - timedelta(days=30)

    if not end_date:
        end_date = datetime.utcnow()

    stmt = (
        select(RiskHistory)
        .where(
            and_(
                RiskHistory.changed_by_id == user_id,
                RiskHistory.changed_at >= start_date,
                RiskHistory.changed_at <= end_date,
            )
        )
        .order_by(desc(RiskHistory.changed_at))
        .limit(limit)
    )

    result = await db.execute(stmt)
    return result.scalars().all()


async def export_audit_log_to_excel(
    db: AsyncSession,
    filters: Optional[Dict[str, Any]] = None,
) -> bytes:
    """
    Export audit trail to Excel for forensic reports.

    Generates a formatted Excel file with audit columns:
      - Risk ID, Timestamp, User, Action
      - Previous Values, New Values (as JSON)
      - IP Address, User-Agent

    Args:
        db: SQLAlchemy async session
        filters: Optional dict with keys: risk_id, user_id, action, date_from, date_to

    Returns:
        Excel file as bytes

    Note:
        This function requires openpyxl to be installed.
        It is included in requirements.txt.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl required for Excel export. Install via: pip install openpyxl")

    # Build query
    query = select(RiskHistory)

    if filters:
        conditions = []

        if "risk_id" in filters:
            conditions.append(RiskHistory.risk_id == filters["risk_id"])

        if "user_id" in filters:
            conditions.append(RiskHistory.changed_by_id == filters["user_id"])

        if "action" in filters:
            conditions.append(RiskHistory.action == filters["action"])

        if "date_from" in filters:
            conditions.append(RiskHistory.changed_at >= filters["date_from"])

        if "date_to" in filters:
            conditions.append(RiskHistory.changed_at <= filters["date_to"])

        if conditions:
            query = query.where(and_(*conditions))

    query = query.order_by(desc(RiskHistory.changed_at))

    result = await db.execute(query)
    entries = result.scalars().all()

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Trail"

    # Headers
    headers = [
        "Risk ID",
        "Timestamp",
        "User ID",
        "Action",
        "Previous Values",
        "New Values",
        "IP Address",
        "User-Agent",
    ]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    for row_num, entry in enumerate(entries, 2):
        ws.cell(row=row_num, column=1).value = str(entry.risk_id)
        ws.cell(row=row_num, column=2).value = entry.changed_at.isoformat()
        ws.cell(row=row_num, column=3).value = str(entry.changed_by_id)
        ws.cell(row=row_num, column=4).value = entry.action
        import json
        ws.cell(row=row_num, column=5).value = json.dumps(entry.previous_values or {}, ensure_ascii=False)
        ws.cell(row=row_num, column=6).value = json.dumps(entry.new_values or {}, ensure_ascii=False)
        ws.cell(row=row_num, column=7).value = entry.ip_address or "-"
        ws.cell(row=row_num, column=8).value = entry.user_agent or "-"

    # Adjust column widths
    for col_num in range(1, len(headers) + 1):
        column_letter = get_column_letter(col_num)
        ws.column_dimensions[column_letter].width = 25

    # Convert to bytes
    import io

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return output.getvalue()
