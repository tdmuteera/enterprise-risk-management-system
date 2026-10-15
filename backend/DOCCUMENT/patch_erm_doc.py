"""
Patch ERM_Prototype_Documentation.docx:
  - Replace [Figure N: ... Insert here] placeholders with PNG images
  - Insert representative code listings after section 5.4 (one-time)
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

DOC_PATH = Path(__file__).resolve().parent / "ERM_Prototype_Documentation.docx"
FIG_DIR = Path(__file__).resolve().parent / "figure_assets"
BACK_ROOT = Path(__file__).resolve().parents[1]

MARKER_ALREADY_PATCHED = "5.4.1 Representative implementation excerpts"


def insert_paragraph_after(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style is not None:
        new_para.style = style
    return new_para


def clear_paragraph_keep_props(paragraph: Paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag.endswith("}pPr"):
            continue
        p.remove(child)


def replace_placeholder_with_figure(paragraph: Paragraph, fig_num: int) -> bool:
    path = FIG_DIR / f"figure{fig_num:02d}.png"
    if not path.exists():
        print(f"Missing image: {path}")
        return False
    clear_paragraph_keep_props(paragraph)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.25))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def figure_num_from_text(text: str) -> int | None:
    t = text.strip().replace("\u2013", "-").replace("\u2014", "-")
    if "insert here" not in t.lower():
        return None
    m = re.search(r"\[Figure\s*(\d+)\s*:", t, re.I)
    if not m:
        return None
    return int(m.group(1))


def read_lines(rel: str, start: int, end_inclusive: int) -> str:
    fp = BACK_ROOT / rel
    lines = fp.read_text(encoding="utf-8").splitlines()
    chunk = lines[start - 1 : end_inclusive]
    return "\n".join(chunk)


def apply_figure_replacements(doc: Document) -> int:
    n = 0
    for para in doc.paragraphs:
        fn = figure_num_from_text(para.text)
        if fn is None:
            continue
        if replace_placeholder_with_figure(para, fn):
            n += 1
    return n


def doc_full_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def add_code_listings_after_54(doc: Document) -> None:
    if MARKER_ALREADY_PATCHED in doc_full_text(doc):
        print("Code listings already present; skipping insertion.")
        return

    anchor: Paragraph | None = None
    for para in doc.paragraphs:
        if "TanStack Query" in para.text and "query invalidation" in para.text:
            anchor = para
            break
    if anchor is None:
        print("Anchor paragraph for code listings not found; skipping.")
        return

    snippets: list[tuple[str, str]] = [
        (
            "Listing 1 — JWT authentication dependency (excerpt from auth/dependencies.py)",
            read_lines("auth/dependencies.py", 1, 85),
        ),
        (
            "Listing 2 — Multi-factor risk scoring utilities (excerpt from utils/risk_scorer.py)",
            read_lines("utils/risk_scorer.py", 17, 71),
        ),
        (
            "Listing 3 — APScheduler background jobs (excerpt from scheduler.py)",
            read_lines("scheduler.py", 1, 96),
        ),
        (
            "Listing 4 — Audit trail write on score change (excerpt from services/scoring_service.py)",
            read_lines("services/scoring_service.py", 85, 132),
        ),
    ]

    cur = anchor
    cur = insert_paragraph_after(cur, "")
    h = insert_paragraph_after(cur, MARKER_ALREADY_PATCHED)
    try:
        h.style = doc.styles["Heading 3"]
    except KeyError:
        h.style = doc.styles["Heading 2"]
    cur = insert_paragraph_after(h, "")

    for title, code in snippets:
        tpara = insert_paragraph_after(cur, title)
        for run in tpara.runs:
            run.bold = True
            run.font.size = Pt(10)
        cur = tpara
        cpara = insert_paragraph_after(cur, code)
        cpara.paragraph_format.left_indent = Inches(0.2)
        cpara.paragraph_format.space_after = Pt(6)
        for run in cpara.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        cur = cpara
        cur = insert_paragraph_after(cur, "")

    print("Inserted code listings after section 5.4.")


def main() -> None:
    if not DOC_PATH.exists():
        raise SystemExit(f"Document not found: {DOC_PATH}")
    bak = DOC_PATH.with_suffix(".pre_figures_backup.docx")
    if not bak.exists():
        shutil.copy2(DOC_PATH, bak)
        print(f"Backup saved: {bak}")

    doc = Document(str(DOC_PATH))
    fig_count = apply_figure_replacements(doc)
    print(f"Inserted {fig_count} figure(s).")
    add_code_listings_after_54(doc)
    doc.save(str(DOC_PATH))
    print(f"Saved: {DOC_PATH}")


if __name__ == "__main__":
    main()
